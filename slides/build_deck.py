# -*- coding: utf-8 -*-
"""Build the internship report deck (Trần Anh Khoa, TGSOFT, Printuz project)."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE / "images"
OUT = HERE / "Bao-cao-thuc-tap-Printuz.pptx"

TEAL = RGBColor(0x4F, 0xA6, 0xA6)
TEAL_DARK = RGBColor(0x2F, 0x7F, 0x80)
INK = RGBColor(0x20, 0x24, 0x2B)
MUTED = RGBColor(0x5B, 0x63, 0x70)
SOFT = RGBColor(0xEA, 0xF3, 0xF3)
SOFT_2 = RGBColor(0xF4, 0xF6, 0xF7)
LINE = RGBColor(0xD9, 0xDE, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PEACH = RGBColor(0xF3, 0xB4, 0x89)

PHASE_TINT = [
    RGBColor(0xE3, 0xF1, 0xF1),
    RGBColor(0xFF, 0xFF, 0xFF),
    RGBColor(0xFD, 0xF0, 0xE6),
    RGBColor(0xF1, 0xF2, 0xF6),
]

FONT = "Calibri"
W = Inches(13.333)
H = Inches(7.5)


def I(v):
    return Inches(v)


def style_run(run, size, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT


def set_bullet(p, indent_in=0.28, color=TEAL):
    pPr = p._p.get_or_add_pPr()
    indent = int(I(indent_in))
    pPr.set("marL", str(indent))
    pPr.set("indent", str(-indent))
    for tag in ("a:buClr", "a:buFont", "a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": str(color)})
    clr.append(srgb)
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, margin=0.05):
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    m = I(margin)
    tf.margin_left = tf.margin_right = m
    tf.margin_top = tf.margin_bottom = m
    return tf


def add_text(slide, text, x, y, w, h, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    tf = textbox(slide, x, y, w, h, anchor)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        style_run(p.add_run(), size, bold, color)
        p.runs[0].text = line
    return tf


def write_items(tf, items, size=16, color=INK, space=6, bullet=True, first=True, lead_color=None):
    """items: str or (lead, rest). Lead is bold."""
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if (first and i == 0) else tf.add_paragraph()
        p.space_after = Pt(space)
        if bullet:
            set_bullet(p)
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run()
            r.text = lead
            style_run(r, size, True, lead_color or color)
            if rest:
                r2 = p.add_run()
                r2.text = rest
                style_run(r2, size, False, color)
        else:
            r = p.add_run()
            r.text = item
            style_run(r, size, False, color)


def bullets(slide, items, x, y, w, h, size=16, color=INK, space=6, bullet=True):
    tf = textbox(slide, x, y, w, h)
    write_items(tf, items, size, color, space, bullet)
    return tf


def rect(slide, x, y, w, h, fill=SOFT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def card(slide, x, y, w, h, kicker, title, items, size=15, fill=SOFT, bullet=True, title_size=19):
    rect(slide, x, y, w, h, fill)
    tf = textbox(slide, x + 0.2, y + 0.16, w - 0.4, h - 0.3)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = kicker
    style_run(r, 12, True, TEAL_DARK)
    if title:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(3)
        p2.space_after = Pt(6)
        r2 = p2.add_run()
        r2.text = title
        style_run(r2, title_size, True, INK)
    else:
        p.space_after = Pt(8)
    write_items(tf, items, size, INK, 5, bullet, first=False)
    return tf


def band(slide, text, y=5.95):
    rect(slide, 0.5, y, 12.33, 0.85, TEAL)
    add_text(slide, text, 0.75, y, 11.9, 0.85, 16, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)


def chrome(slide, section, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, I(0.09))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    if section:
        add_text(slide, section.upper(), 0.5, 0.26, 9, 0.3, 12, True, TEAL_DARK)
    add_text(slide, title, 0.5, 0.5, 12.3, 0.62, 28, True, INK)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(0.5), I(7.05), I(12.83), I(7.05))
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)
    add_text(slide, "Trần Anh Khoa  ·  Báo cáo thực tập tốt nghiệp  ·  Website Printuz", 0.5, 7.1, 9, 0.3, 11,
             False, MUTED)


def picture(slide, path, x, y, max_w, max_h, border=True):
    with Image.open(path) as im:
        iw, ih = im.size
    aspect = iw / ih
    if aspect > max_w / max_h:
        w, h = max_w, max_w / aspect
    else:
        w, h = max_h * aspect, max_h
    left = x + (max_w - w) / 2
    top = y + (max_h - h) / 2
    slide.shapes.add_picture(str(path), I(left), I(top), I(w), I(h))
    if border:
        f = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(left), I(top), I(w), I(h))
        f.fill.background()
        f.line.color.rgb = LINE
        f.line.width = Pt(1)
        f.shadow.inherit = False
    return left, top, w, h


def figure(slide, path, x, y, max_w, max_h, caption):
    left, top, w, h = picture(slide, path, x, y, max_w, max_h)
    add_text(slide, caption, x, top + h + 0.06, max_w, 0.3, 12, False, MUTED, PP_ALIGN.CENTER)


def table(slide, x, y, col_w, rows, header_size=13, body_size=12, row_h=0.37, row_fills=None):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, I(x), I(y), I(sum(col_w)), I(row_h * n_rows))
    tbl = shape.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = I(cw)
    for i in range(n_rows):
        tbl.rows[i].height = I(row_h)
        for j in range(n_cols):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = I(0.08)
            cell.margin_top = cell.margin_bottom = I(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = rows[i][j]
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = TEAL
                style_run(r, header_size, True, WHITE)
            else:
                cell.fill.fore_color.rgb = row_fills[i - 1] if row_fills else (SOFT_2 if i % 2 == 0 else WHITE)
                style_run(r, body_size, j == 0, INK)
    return tbl


def chevrons(slide, labels, x, y, w, h, size=13):
    n = len(labels)
    overlap = 0.12
    step = (w + overlap * (n - 1)) / n
    for i, lab in enumerate(labels):
        shp_type = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        s = slide.shapes.add_shape(shp_type, I(x + i * (step - overlap)), I(y), I(step), I(h))
        s.fill.solid()
        s.fill.fore_color.rgb = TEAL if i % 2 == 0 else TEAL_DARK
        s.line.fill.background()
        s.shadow.inherit = False
        tf = s.text_frame
        tf.word_wrap = True
        tf.margin_left = I(0.25 if i else 0.12)
        tf.margin_right = I(0.12)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = lab
        style_run(r, size, True, WHITE)


def stat(slide, x, y, w, h, big, small):
    rect(slide, x, y, w, h, SOFT)
    add_text(slide, big, x, y + 0.15, w, 0.6, 30, True, TEAL_DARK, PP_ALIGN.CENTER)
    add_text(slide, small, x + 0.1, y + 0.8, w - 0.2, h - 0.85, 14, False, INK, PP_ALIGN.CENTER)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def line(slide, x1, y1, x2, y2, color=TEAL, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def org_box(slide, x, y, w, h, title, sub, fill=TEAL, color=WHITE, dashed=False):
    s = rect(slide, x, y, w, h, fill, None if not dashed else TEAL_DARK, radius=0.12)
    if dashed:
        s.line.dash_style = 4
        s.line.width = Pt(1.5)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    style_run(r, 15, True, color)
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        style_run(r2, 11, False, color)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Báo cáo thực tập tốt nghiệp — Website Printuz"
    prs.core_properties.author = "Trần Anh Khoa"
    blank = prs.slide_layouts[6]

    def new(section, title):
        s = prs.slides.add_slide(blank)
        chrome(s, section, title)
        return s

    # ---------------------------------------------------------------- 1 Cover
    s = prs.slides.add_slide(blank)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, I(0.09))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    add_text(s, "TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH\nKHOA CÔNG NGHỆ THÔNG TIN", 0.6, 0.45, 7, 0.7, 13,
             True, TEAL_DARK)
    add_text(s, "BÁO CÁO THỰC TẬP TỐT NGHIỆP", 0.6, 1.45, 7, 0.4, 16, True, MUTED)
    add_text(s, "Xây dựng giao diện website\ngiới thiệu dịch vụ cho Printuz LTD", 0.6, 1.85, 6.6, 1.4, 32, True,
             INK)
    tf = textbox(s, 0.6, 3.6, 6.3, 2.9)
    write_items(
        tf,
        [
            ("Sinh viên: ", "Trần Anh Khoa  ·  MSSV 2001230402  ·  Lớp 14DHTH08"),
            ("Giảng viên hướng dẫn: ", "ThS. Nguyễn Thị Định"),
            ("Đơn vị thực tập: ", "Công ty TNHH Công nghệ TGSOFT"),
            ("Người hướng dẫn tại công ty: ", "Anh Trần Thế Phú"),
            ("Thời gian: ", "06/07/2026 – 27/09/2026 (12 tuần)"),
        ],
        16,
        INK,
        8,
        bullet=False,
    )
    add_text(s, "TP. Hồ Chí Minh, tháng 9 năm 2026", 0.6, 6.55, 6, 0.35, 13, False, MUTED)
    picture(s, SHOTS / "printuz-about.png", 7.35, 1.2, 5.5, 5.2)
    notes(
        s,
        """
Em kính chào quý thầy cô.
Em là Trần Anh Khoa, mã số sinh viên 2001230402, lớp 14DHTH08, Khoa Công nghệ Thông tin.
Em thực tập tại Công ty TNHH Công nghệ TGSOFT trong 12 tuần, từ ngày 06/07 đến ngày 27/09/2026.
Giảng viên hướng dẫn của em là cô Nguyễn Thị Định. Người hướng dẫn trực tiếp tại công ty là anh Trần Thế Phú.
Đề tài em thực hiện là xây dựng giao diện website giới thiệu dịch vụ cho Printuz LTD.
Hình bên phải là trang giới thiệu của website em đã làm.
""",
    )

    # ---------------------------------------------------------------- 2 Agenda
    s = new("", "Nội dung trình bày")
    agenda = [
        ("01", "Giới thiệu đơn vị thực tập", "Thông tin, lĩnh vực hoạt động, cơ cấu tổ chức của TGSOFT"),
        ("02", "Mục tiêu, nội dung, ý nghĩa", "Công việc thực tập hướng tới điều gì và để làm gì"),
        ("03", "Vị trí công việc thực tập", "Vị trí Frontend, yêu cầu và nhiệm vụ được giao"),
        ("04", "Kế hoạch và tiến độ", "Bốn giai đoạn và nhật ký 12 tuần"),
        ("05", "Phân tích và công nghệ", "Giao diện, trải nghiệm người dùng, yêu cầu dự án"),
        ("06", "Kết quả thực hiện", "Giao diện website Printuz đã hoàn thành"),
        ("07", "Đánh giá và kiến nghị", "Ưu điểm, hạn chế, bài học, đề xuất"),
        ("08", "Kết luận", ""),
    ]
    for i, (num, head, sub) in enumerate(agenda):
        col, row = divmod(i, 4)
        x = 0.6 + col * 6.25
        y = 1.45 + row * 1.33
        rect(s, x, y, 5.95, 1.12, SOFT)
        add_text(s, num, x + 0.2, y + 0.18, 0.9, 0.75, 30, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, head, x + 1.15, y + 0.17, 4.6, 0.42, 19, True, INK)
        if sub:
            add_text(s, sub, x + 1.15, y + 0.6, 4.6, 0.4, 13, False, MUTED)
    notes(
        s,
        """
Bài trình bày của em gồm tám phần, đi theo bố cục của báo cáo.
Đầu tiên em giới thiệu đơn vị thực tập là Công ty TGSOFT.
Tiếp theo là mục tiêu, nội dung và ý nghĩa của công việc thực tập.
Phần ba là vị trí công việc em đảm nhận, và phần bốn là kế hoạch thực tập trong 12 tuần.
Sau đó em trình bày phần phân tích, công nghệ sử dụng và kết quả website.
Cuối cùng là đánh giá, kiến nghị và kết luận.
""",
    )

    # ---------------------------------------------------------------- 3 Company info
    s = new("01 · Giới thiệu đơn vị thực tập", "Công ty TNHH Công nghệ TGSOFT")
    rect(s, 0.5, 1.35, 6.15, 5.45, SOFT)
    add_text(s, "THÔNG TIN CHUNG", 0.75, 1.5, 5.6, 0.3, 12, True, TEAL_DARK)
    rows = [
        ("Tên công ty", "Công ty TNHH Công nghệ TGSOFT"),
        ("Mã số thuế", "0315310284"),
        ("Đại diện pháp luật", "Ông Trần Thế Văn"),
        ("Địa chỉ", "Tầng 14, Tòa nhà HM Town, 412 Nguyễn Thị Minh Khai, phường Bàn Cờ, TP. Hồ Chí Minh"),
        ("Email", "admin@tgsoft.vn"),
        ("Hotline", "0919 33 2220"),
        ("Người hướng dẫn", "Anh Trần Thế Phú"),
    ]
    yy = 1.88
    for k, v in rows:
        h = 0.78 if k == "Địa chỉ" else 0.55
        add_text(s, k, 0.75, yy, 1.85, h, 14, True, MUTED)
        add_text(s, v, 2.6, yy, 3.9, h, 15, False, INK)
        yy += h + 0.08
    card(
        s,
        6.9,
        1.35,
        5.93,
        5.45,
        "SƠ LƯỢC HÌNH THÀNH VÀ PHÁT TRIỂN",
        "Thành lập năm 2014",
        [
            "Do những người Việt Nam trẻ tuổi, sáng tạo, nhiệt huyết với công nghệ thành lập.",
            "Hoạt động trong lĩnh vực công nghệ thông tin: sản phẩm, dịch vụ và giải pháp phần mềm, Internet.",
            "Từng bước phát triển dịch vụ lập trình, xây dựng và quản trị website, phát triển ứng dụng.",
            "Chú trọng cả quá trình làm sản phẩm: phân tích yêu cầu, thiết kế, lập trình, kiểm tra, bảo trì và hỗ trợ sau bàn giao.",
            "Cung cấp thêm quảng cáo trực tuyến, SEO, hosting, tên miền và quản trị website.",
        ],
        15,
        fill=SOFT_2,
    )
    notes(
        s,
        """
Đơn vị em thực tập là Công ty TNHH Công nghệ TGSOFT, mã số thuế 0315310284.
Người đại diện theo pháp luật là ông Trần Thế Văn. Công ty đặt tại tầng 14, tòa nhà HM Town, 412 Nguyễn Thị Minh Khai, phường Bàn Cờ, Thành phố Hồ Chí Minh.
Người hướng dẫn trực tiếp của em tại công ty là anh Trần Thế Phú.
Theo giới thiệu trên website của công ty, TGSOFT được thành lập năm 2014 bởi những người Việt Nam trẻ, có nhiệt huyết với công nghệ.
Công ty phát triển dần các dịch vụ lập trình, xây dựng và quản trị website, phát triển ứng dụng, và các dịch vụ hỗ trợ kinh doanh trên Internet.
Điểm em thấy rõ là công ty chú trọng cả quá trình làm một sản phẩm, từ phân tích yêu cầu đến bảo trì sau khi bàn giao.
""",
    )

    # ---------------------------------------------------------------- 4 Fields + process
    s = new("01 · Giới thiệu đơn vị thực tập", "Lĩnh vực hoạt động và quy trình làm sản phẩm")
    fields = [
        ("Lập trình website", "Xây dựng website theo yêu cầu khách hàng"),
        ("Android, iOS, WebApp", "Ứng dụng di động và ứng dụng chạy trên web"),
        ("Website 60s", "Website triển khai nhanh cho nhu cầu cơ bản"),
        ("Quản trị website, fanpage", "Cập nhật, duy trì nội dung"),
        ("Google Ads và SEO", "Quảng bá website trên công cụ tìm kiếm"),
        ("Quảng cáo Facebook", "Triển khai quảng cáo trên Facebook"),
        ("Phát triển kênh", "Tăng theo dõi fanpage, kênh YouTube"),
        ("Hosting và tên miền", "Lưu trữ website và tên miền"),
    ]
    for i, (head, sub) in enumerate(fields):
        col, row = i % 4, i // 4
        x = 0.5 + col * 3.1
        y = 1.4 + row * 1.5
        rect(s, x, y, 2.92, 1.32, SOFT)
        add_text(s, head, x + 0.15, y + 0.14, 2.65, 0.4, 16, True, INK)
        add_text(s, sub, x + 0.15, y + 0.58, 2.65, 0.7, 13, False, MUTED)
    add_text(s, "QUY TRÌNH PHÁT TRIỂN PHẦN MỀM VÀ WEBSITE", 0.5, 4.55, 9, 0.3, 12, True, TEAL_DARK)
    chevrons(s, ["Phân tích", "Thiết kế", "Lập trình", "Kiểm tra", "Bảo trì"], 0.5, 4.95, 12.33, 0.85, 17)
    add_text(
        s,
        "Công việc chia theo từng giai đoạn, nên dễ kiểm tra và điều chỉnh sản phẩm trước và sau khi bàn giao cho khách hàng.",
        0.5,
        6.0,
        12.3,
        0.6,
        15,
        False,
        MUTED,
    )
    notes(
        s,
        """
TGSOFT hoạt động chủ yếu trong lĩnh vực công nghệ thông tin, với tám nhóm dịch vụ trên slide.
Lĩnh vực gắn trực tiếp với công việc của em là lập trình website.
Bên cạnh đó công ty còn làm ứng dụng Android, iOS, WebApp, gói Website 60s, quản trị website và fanpage, quảng cáo Google, SEO, quảng cáo Facebook, phát triển kênh, hosting và tên miền.
Với phần mềm và website, công ty làm theo quy trình năm bước: phân tích, thiết kế, lập trình, kiểm tra và bảo trì.
Em cũng làm dự án của mình theo đúng tinh thần này: tìm hiểu và phân tích trước, rồi xây giao diện, kiểm tra, sửa lỗi theo góp ý.
""",
    )

    # ---------------------------------------------------------------- 5 Org chart
    s = new("01 · Giới thiệu đơn vị thực tập", "Tổ chức quản lý và cơ cấu tổ chức")
    rect(s, 0.5, 1.35, 7.0, 5.45, SOFT_2)
    org_box(s, 2.75, 1.6, 2.5, 0.8, "Giám đốc", "Định hướng chung")
    line(s, 4.0, 2.4, 4.0, 2.8)
    line(s, 1.65, 2.8, 6.35, 2.8)
    for cx in (1.65, 4.0, 6.35):
        line(s, cx, 2.8, cx, 3.1)
    org_box(s, 0.75, 3.1, 1.8, 0.8, "Trưởng nhóm", "Phân công kỹ thuật")
    org_box(s, 3.1, 3.1, 1.8, 0.8, "Kế toán", "Tài chính")
    org_box(s, 5.45, 3.1, 1.8, 0.8, "Kinh doanh", "Khách hàng, tư vấn")
    line(s, 1.65, 3.9, 1.65, 4.3)
    org_box(s, 0.75, 4.3, 2.5, 0.8, "Nhân viên kỹ thuật", "Lập trình, bảo trì", fill=TEAL_DARK)
    line(s, 1.65, 5.1, 1.65, 5.45, TEAL_DARK)
    org_box(s, 0.75, 5.45, 3.2, 0.95, "Thực tập sinh Frontend", "Vị trí của em, trong nhóm kỹ thuật", fill=WHITE,
            color=TEAL_DARK, dashed=True)
    tf = textbox(s, 7.8, 1.35, 5.0, 5.5)
    write_items(
        tf,
        [
            ("Giám đốc: ", "quản lý và đưa ra định hướng chung cho công ty."),
            ("Trưởng nhóm: ", "quản lý, phân công công việc cho nhóm kỹ thuật."),
            ("Nhân viên kỹ thuật: ", "lập trình, phát triển và bảo trì sản phẩm."),
            ("Kế toán: ", "các công việc tài chính, kế toán."),
            ("Kinh doanh: ", "làm việc với khách hàng, tư vấn, hỗ trợ kinh doanh."),
            ("Nguồn lực: ", "máy tính, phần mềm và công cụ cho lập trình, kiểm tra, triển khai, dùng theo từng dự án."),
            ("Đặc điểm: ", "cơ cấu gọn, phối hợp giữa kỹ thuật và các bộ phận khác thuận tiện."),
        ],
        15,
        INK,
        7,
    )
    notes(
        s,
        """
TGSOFT tổ chức nhân sự phù hợp với quy mô của một công ty nhỏ.
Đứng đầu là Giám đốc, giữ vai trò quản lý và định hướng chung.
Bên dưới có trưởng nhóm, kế toán và kinh doanh. Trưởng nhóm quản lý và phân công công việc cho nhân viên kỹ thuật, là những người lập trình, phát triển và bảo trì sản phẩm.
Em làm thực tập sinh Frontend trong nhóm kỹ thuật, nhận việc và được hướng dẫn trực tiếp bởi anh Trần Thế Phú.
Vì cơ cấu gọn nên việc trao đổi giữa các vị trí khá thuận tiện.
Do công ty nhỏ và có một số thông tin nhạy cảm, công ty chỉ cung cấp thông tin sơ lược, nên phần này em trình bày ở mức tổng quát.
""",
    )

    # ---------------------------------------------------------------- 6 Objectives
    s = new("02 · Mục tiêu, nội dung, ý nghĩa", "Mục tiêu của công việc thực tập")
    rect(s, 0.5, 1.35, 12.33, 1.15, TEAL)
    add_text(s, "MỤC TIÊU CHUNG", 0.75, 1.45, 4, 0.3, 12, True, WHITE)
    add_text(
        s,
        "Vận dụng kiến thức đã học vào một dự án thực tế tại doanh nghiệp: hoàn thành giao diện website giới thiệu dịch vụ fulfillment cho Printuz LTD.",
        0.75,
        1.75,
        11.9,
        0.7,
        17,
        True,
        WHITE,
    )
    objectives = [
        ("Quy trình làm việc", "Làm quen quy trình thực tế: nhận yêu cầu, thực hiện, kiểm tra, báo cáo với người hướng dẫn."),
        ("Giao diện đa trang", "Xây dựng các trang của website Printuz với bố cục rõ ràng, thống nhất màu sắc và thương hiệu."),
        ("Đáp ứng thiết bị", "Bảo đảm hiển thị tốt trên máy tính, máy tính bảng và điện thoại."),
        ("Tổ chức dữ liệu", "Sắp xếp dữ liệu sản phẩm và bài viết để website dễ mở rộng về sau."),
        ("Kỹ năng bản thân", "Rèn tự học, làm việc nhóm, giao tiếp và tiếp nhận góp ý."),
    ]
    for i, (head, body) in enumerate(objectives):
        x = 0.5 + i * 2.5
        rect(s, x, 2.75, 2.33, 3.0, SOFT)
        add_text(s, f"0{i + 1}", x + 0.18, 2.88, 1, 0.5, 24, True, TEAL)
        add_text(s, head, x + 0.18, 3.4, 2.0, 0.45, 17, True, INK)
        add_text(s, body, x + 0.18, 3.9, 2.0, 1.8, 14, False, MUTED)
    rect(s, 0.5, 5.95, 12.33, 0.85, WHITE, line=LINE)
    tf = textbox(s, 0.75, 6.07, 11.9, 0.65)
    write_items(
        tf,
        [("Kết quả mong đợi: ", "bản demo website Printuz hoàn chỉnh, hiển thị tốt trên nhiều thiết bị, cùng bài báo cáo thực tập tổng hợp quá trình làm việc.")],
        15,
        INK,
        0,
        False,
    )
    notes(
        s,
        """
Mục tiêu chung của kỳ thực tập là vận dụng kiến thức đã học vào một dự án thật tại doanh nghiệp. Cụ thể là hoàn thành giao diện website giới thiệu dịch vụ fulfillment cho Printuz.
Từ đó em đặt ra năm mục tiêu cụ thể.
Thứ nhất, làm quen quy trình làm việc thực tế: nhận yêu cầu, thực hiện, kiểm tra và báo cáo với người hướng dẫn.
Thứ hai, xây dựng giao diện đa trang cho Printuz, bố cục rõ ràng và thống nhất với thương hiệu.
Thứ ba, website phải hiển thị tốt trên máy tính, máy tính bảng và điện thoại.
Thứ tư, dữ liệu sản phẩm và bài viết được tổ chức để có thể mở rộng sau này.
Thứ năm là rèn các kỹ năng của bản thân như tự học, làm việc nhóm và tiếp nhận góp ý.
""",
    )

    # ---------------------------------------------------------------- 7 Content
    s = new("02 · Mục tiêu, nội dung, ý nghĩa", "Nội dung công việc thực tập")
    card(
        s,
        0.5,
        1.35,
        6.05,
        2.45,
        "ĐỀ TÀI",
        None,
        [
            "“Thiết kế và xây dựng giao diện website fulfillment Printuz tại Công ty Công nghệ TGSOFT”.",
            "Tham khảo cấu trúc website Printuz, xây dựng phiên bản giao diện mới hiện đại, thống nhất, thân thiện hơn.",
        ],
        15,
        bullet=True,
    )
    card(
        s,
        0.5,
        4.0,
        6.05,
        2.8,
        "ĐỐI TƯỢNG NGƯỜI DÙNG",
        None,
        [
            "Cá nhân, doanh nghiệp kinh doanh thương mại điện tử.",
            "Có nhu cầu in theo yêu cầu (POD), tìm nguồn hàng, dropshipping và fulfillment.",
            "Vì vậy giao diện phải nói rõ dịch vụ, tạo cảm giác chuyên nghiệp, giúp tìm thông tin nhanh.",
        ],
        15,
        fill=SOFT_2,
    )
    add_text(s, "CÁC HẠNG MỤC THỰC HIỆN", 6.85, 1.35, 6, 0.3, 12, True, TEAL_DARK)
    steps = [
        ("Khảo sát", "cấu trúc website Printuz, các trang chính, danh mục sản phẩm, thành phần giao diện."),
        ("Thành phần dùng chung", "header, menu điều hướng, footer, ô tìm kiếm, nút về đầu trang."),
        ("Các trang", "trang chủ, giới thiệu, dịch vụ, sản phẩm, tin tức, câu hỏi thường gặp, liên hệ, tài khoản, hướng dẫn extension."),
        ("Responsive", "chỉnh hiển thị trên máy tính, máy tính bảng, điện thoại bằng Bootstrap."),
        ("Kiểm tra, sửa lỗi", "bố cục, font chữ, khoảng cách, hình ảnh, tương thích trình duyệt."),
    ]
    for i, (head, body) in enumerate(steps):
        y = 1.75 + i * 1.0
        rect(s, 6.85, y, 0.55, 0.55, TEAL, shape=MSO_SHAPE.OVAL)
        add_text(s, str(i + 1), 6.85, y, 0.55, 0.55, 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        tf = textbox(s, 7.55, y - 0.05, 5.3, 0.95)
        write_items(tf, [(head + ": ", body)], 14, INK, 0, bullet=False)
    notes(
        s,
        """
Đề tài của em là thiết kế và xây dựng giao diện website fulfillment Printuz tại Công ty TGSOFT.
Em tham khảo cấu trúc website Printuz hiện có, rồi xây dựng phiên bản giao diện mới theo hướng hiện đại, thống nhất và thân thiện hơn.
Người dùng của website là các cá nhân, doanh nghiệp bán hàng trực tuyến cần in theo yêu cầu, tìm nguồn hàng, dropshipping hoặc fulfillment. Vì vậy giao diện phải nói rõ dịch vụ và giúp họ tìm thông tin nhanh.
Công việc gồm năm hạng mục: khảo sát website, làm các thành phần dùng chung, xây dựng từng trang, chỉnh hiển thị trên nhiều thiết bị, và cuối cùng là kiểm tra, sửa lỗi.
""",
    )

    # ---------------------------------------------------------------- 8 Significance
    s = new("02 · Mục tiêu, nội dung, ý nghĩa", "Ý nghĩa của công việc thực tập")
    sig = [
        (
            "ĐỐI VỚI BẢN THÂN",
            "Áp dụng kiến thức vào dự án thật",
            [
                "Dùng HTML, CSS, Bootstrap, JavaScript cho một website hoàn chỉnh.",
                "Hiểu tác phong và quy trình làm việc tại doanh nghiệp.",
                "Có kinh nghiệm thực tế, tự tin hơn cho nghề nghiệp sau tốt nghiệp.",
            ],
        ),
        (
            "ĐỐI VỚI DOANH NGHIỆP",
            "Giao diện nền tảng cho Printuz",
            [
                "Phiên bản giao diện mới, hiện đại, thống nhất thương hiệu.",
                "Dùng làm bản demo để trình bày và định hướng phát triển.",
                "Là nền để tích hợp dữ liệu động và chức năng nghiệp vụ sau này.",
            ],
        ),
        (
            "ĐỐI VỚI NGƯỜI DÙNG",
            "Tìm thông tin nhanh và dễ hơn",
            [
                "Tiếp cận nhanh dịch vụ, sản phẩm, hướng dẫn, chính sách, tin tức.",
                "Dùng được trên máy tính và điện thoại.",
                "Giao diện chuyên nghiệp giúp tăng nhận diện và sự tin cậy.",
            ],
        ),
    ]
    for i, (k, t, items) in enumerate(sig):
        card(s, 0.5 + i * 4.18, 1.4, 3.97, 4.3, k, t, items, 16, fill=SOFT if i != 1 else SOFT_2)
    band(s, "Kỳ thực tập vừa đáp ứng yêu cầu của học phần, vừa tạo ra một sản phẩm có giá trị thực tế cho doanh nghiệp.")
    notes(
        s,
        """
Công việc thực tập có ý nghĩa ở ba phía.
Với bản thân em, đây là lần đầu áp dụng HTML, CSS, Bootstrap và JavaScript vào một website hoàn chỉnh, trong môi trường doanh nghiệp thật. Em hiểu thêm tác phong làm việc và có kinh nghiệm để tự tin hơn sau khi tốt nghiệp.
Với doanh nghiệp, Printuz có một phiên bản giao diện mới, thống nhất thương hiệu. Bản này dùng làm demo để trình bày, đồng thời là nền để tích hợp dữ liệu động và các chức năng nghiệp vụ về sau.
Với người dùng, họ tìm được dịch vụ, sản phẩm và hướng dẫn nhanh hơn, dùng được trên cả điện thoại, và có cảm giác tin cậy hơn khi xem website.
""",
    )

    # ---------------------------------------------------------------- 9 Position info
    s = new("03 · Vị trí công việc thực tập", "Thông tin vị trí thực tập")
    info = [
        ["Hạng mục", "Thông tin"],
        ["Vị trí", "Thực tập sinh lập trình giao diện website (Frontend)"],
        ["Bộ phận", "Nhóm kỹ thuật"],
        ["Người hướng dẫn", "Anh Trần Thế Phú"],
        ["Thời gian", "12 tuần, từ 06/07/2026 đến 27/09/2026"],
        ["Dự án được giao", "Giao diện website giới thiệu dịch vụ Printuz LTD"],
        ["Hình thức làm việc", "Làm trong nhóm, có giám sát và hướng dẫn của mentor"],
    ]
    table(s, 0.5, 1.4, [2.2, 4.05], info, 14, 14, 0.62)
    card(
        s,
        7.0,
        1.4,
        5.83,
        2.55,
        "GIỚI THIỆU CHUNG VỀ VỊ TRÍ",
        None,
        [
            "Xây dựng phần giao diện mà người dùng trực tiếp nhìn thấy và tương tác.",
            "Chuyển yêu cầu và thiết kế thành trang web có cấu trúc rõ ràng, dễ dùng, phù hợp nhiều thiết bị.",
        ],
        17,
    )
    card(
        s,
        7.0,
        4.15,
        5.83,
        2.6,
        "ĐẶC ĐIỂM CÔNG VIỆC",
        None,
        [
            "Vừa cần kiến thức lập trình web, vừa cần khả năng quan sát và xử lý giao diện.",
            "Làm theo vòng: tìm hiểu yêu cầu, xây dựng giao diện, kiểm tra hoạt động, sửa lỗi phát sinh.",
        ],
        17,
        fill=SOFT_2,
    )
    notes(
        s,
        """
Trong thời gian thực tập, em đảm nhận vị trí thực tập sinh lập trình giao diện website, tức Frontend, thuộc nhóm kỹ thuật.
Người hướng dẫn trực tiếp là anh Trần Thế Phú. Thời gian là 12 tuần, từ 06/07 đến 27/09/2026.
Dự án được giao là giao diện website giới thiệu dịch vụ cho Printuz LTD.
Vị trí Frontend xây dựng phần người dùng trực tiếp nhìn thấy và tương tác. Nhiệm vụ là chuyển yêu cầu và thiết kế thành các trang web rõ ràng, dễ dùng và phù hợp nhiều thiết bị.
Đặc điểm của công việc này là vừa cần kiến thức lập trình, vừa cần con mắt quan sát giao diện. Em làm theo vòng: tìm hiểu yêu cầu, xây dựng, kiểm tra, rồi sửa lỗi.
""",
    )

    # ---------------------------------------------------------------- 10 Requirements
    s = new("03 · Vị trí công việc thực tập", "Yêu cầu đối với thực tập sinh")
    card(
        s,
        0.5,
        1.4,
        3.95,
        4.3,
        "VỀ KIẾN THỨC",
        "Nền tảng lập trình web",
        [
            "Kiến thức cơ bản về lập trình website.",
            "Nắm được HTML, CSS, Bootstrap và JavaScript.",
            "Hiểu khái niệm cơ bản về giao diện người dùng (UI) và trải nghiệm người dùng (UX).",
        ],
        15,
    )
    card(
        s,
        4.69,
        1.4,
        3.95,
        4.3,
        "VỀ KỸ NĂNG",
        "Làm được và làm có trách nhiệm",
        [
            "Xây dựng được giao diện website.",
            "Biết kiểm tra và sửa lỗi giao diện.",
            "Tự học, tự tìm tài liệu.",
            "Có tinh thần trách nhiệm với công việc được giao.",
        ],
        15,
        fill=SOFT_2,
    )
    card(
        s,
        8.88,
        1.4,
        3.95,
        4.3,
        "CÁCH EM ĐÁP ỨNG",
        "Chuẩn bị trong kỳ thực tập",
        [
            "Tuần 3 dành riêng để tìm hiểu các thành phần và cách dùng Bootstrap.",
            "Tra cứu tài liệu MDN Web Docs và tài liệu chính thức của Bootstrap 5.",
            "Dùng công cụ AI để tự học và xử lý vấn đề phát sinh.",
            "Hỏi lại người hướng dẫn khi yêu cầu chưa rõ.",
        ],
        15,
    )
    band(s, "Yêu cầu chung: tìm hiểu yêu cầu, xây dựng giao diện, kiểm tra hoạt động và điều chỉnh lỗi phát sinh.")
    notes(
        s,
        """
Vị trí Frontend đặt ra yêu cầu về kiến thức và kỹ năng.
Về kiến thức, thực tập sinh cần biết lập trình website cơ bản, nắm HTML, CSS, Bootstrap, JavaScript, và hiểu khái niệm giao diện người dùng, trải nghiệm người dùng.
Về kỹ năng, cần xây dựng được giao diện, biết kiểm tra và sửa lỗi, có khả năng tự học và có trách nhiệm với công việc.
Để đáp ứng, em dành riêng tuần thứ ba để tìm hiểu Bootstrap. Em tra cứu tài liệu MDN và tài liệu Bootstrap, dùng công cụ AI để tự học, và hỏi lại anh Phú khi yêu cầu chưa rõ.
""",
    )

    # ---------------------------------------------------------------- 11 Tasks and workflow
    s = new("03 · Vị trí công việc thực tập", "Nhiệm vụ được giao và cách phối hợp")
    tf = textbox(s, 0.5, 1.4, 6.3, 5.4)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "NHIỆM VỤ ĐƯỢC GIAO"
    style_run(r, 12, True, TEAL_DARK)
    p.space_after = Pt(8)
    write_items(
        tf,
        [
            "Tìm hiểu yêu cầu của giao diện website.",
            "Xây dựng giao diện bằng HTML, CSS và JavaScript.",
            "Điều chỉnh bố cục và hiển thị các thành phần trên website.",
            "Kiểm tra hiển thị trên các kích thước màn hình khác nhau.",
            "Phát hiện và sửa lỗi giao diện.",
            "Cập nhật, hoàn thiện giao diện theo yêu cầu của người hướng dẫn.",
            "Báo cáo tiến độ, trao đổi kết quả công việc.",
        ],
        16,
        INK,
        8,
        first=False,
    )
    rect(s, 7.1, 1.4, 5.73, 5.35, SOFT)
    add_text(s, "CÁCH PHỐI HỢP VỚI NGƯỜI HƯỚNG DẪN", 7.35, 1.55, 5.3, 0.3, 12, True, TEAL_DARK)
    flow = [
        "Nhận nhiệm vụ từ anh Trần Thế Phú",
        "Tìm hiểu yêu cầu",
        "Thực hiện công việc",
        "Tự kiểm tra kết quả",
        "Trao đổi lại khi chưa rõ, điều chỉnh theo góp ý",
        "Báo cáo tiến độ",
    ]
    for i, lab in enumerate(flow):
        y = 2.0 + i * 0.75
        rect(s, 7.35, y, 5.25, 0.58, WHITE if i % 2 else TEAL, radius=0.2)
        add_text(s, lab, 7.45, y, 5.05, 0.58, 15, True, TEAL_DARK if i % 2 else WHITE, PP_ALIGN.CENTER,
                 MSO_ANCHOR.MIDDLE)
        if i < len(flow) - 1:
            line(s, 9.97, y + 0.58, 9.97, y + 0.75, TEAL_DARK, 1.5)
    notes(
        s,
        """
Bên trái là bảy nhiệm vụ em được giao: tìm hiểu yêu cầu, xây dựng giao diện bằng HTML, CSS, JavaScript, chỉnh bố cục, kiểm tra trên nhiều màn hình, sửa lỗi, hoàn thiện theo yêu cầu, và báo cáo tiến độ.
Bên phải là cách em phối hợp với người hướng dẫn.
Em nhận nhiệm vụ từ anh Phú, tìm hiểu yêu cầu, làm, rồi tự kiểm tra kết quả. Chỗ nào chưa rõ em hỏi lại để được hỗ trợ và điều chỉnh, sau đó báo cáo tiến độ.
Cách làm này cho em không gian tự chủ, nhưng vẫn có hỗ trợ khi gặp khó khăn.
""",
    )

    # ---------------------------------------------------------------- 12 Plan (4 phases)
    s = new("04 · Kế hoạch và tiến độ", "Kế hoạch thực tập: bốn giai đoạn")
    phases = [
        ("GIAI ĐOẠN 1", "Tuần 1 – 3", "06/07 – 26/07", "Nghiên cứu, phân tích yêu cầu",
         ["Tìm hiểu lĩnh vực fulfillment, khách hàng mục tiêu.", "Khảo sát bố cục, màu sắc, cấu trúc website Printuz.",
          "Tìm hiểu Bootstrap."],
         "Định hướng thiết kế và danh sách trang cần làm."),
        ("GIAI ĐOẠN 2", "Tuần 4 – 8", "27/07 – 30/08", "Xây dựng giao diện",
         ["Header, menu, footer, trang chủ.", "Catalog, chi tiết sản phẩm, giới thiệu, liên hệ, blog.",
          "Responsive bằng Bootstrap."],
         "Các trang chính của website."),
        ("GIAI ĐOẠN 3", "Tuần 9 – 10", "31/08 – 13/09", "Hoàn thiện và kiểm thử",
         ["Rà soát bố cục, màu, font, khoảng cách, hiệu ứng.", "Sửa lỗi theo góp ý.", "Báo cáo người hướng dẫn."],
         "Bản demo website hoàn chỉnh."),
        ("GIAI ĐOẠN 4", "Tuần 11 – 12", "14/09 – 27/09", "Viết báo cáo",
         ["Tổng hợp nội dung.", "Viết, rà soát báo cáo.", "Hoàn tất theo yêu cầu của trường."],
         "Báo cáo thực tập."),
    ]
    line(s, 0.8, 1.72, 12.5, 1.72, TEAL, 2.5)
    for i, (k, wk, dt, title, items, outp) in enumerate(phases):
        x = 0.5 + i * 3.1
        dot = rect(s, x + 1.33, 1.55, 0.34, 0.34, TEAL if i % 2 == 0 else TEAL_DARK, shape=MSO_SHAPE.OVAL)
        rect(s, x, 2.05, 2.92, 4.75, PHASE_TINT[i] if i != 1 else SOFT, LINE if i == 1 else None)
        add_text(s, k, x + 0.15, 2.15, 2.6, 0.3, 12, True, TEAL_DARK)
        add_text(s, wk, x + 0.15, 2.42, 2.6, 0.4, 20, True, INK)
        add_text(s, dt, x + 0.15, 2.83, 2.6, 0.3, 13, False, MUTED)
        add_text(s, title, x + 0.15, 3.15, 2.65, 0.65, 16, True, TEAL_DARK)
        tf = textbox(s, x + 0.1, 3.8, 2.75, 1.95)
        write_items(tf, items, 13, INK, 3)
        rect(s, x + 0.12, 5.8, 2.68, 0.85, WHITE, LINE, radius=0.1)
        tf = textbox(s, x + 0.2, 5.83, 2.55, 0.8, MSO_ANCHOR.MIDDLE)
        write_items(tf, [("Kết quả: ", outp)], 12, INK, 0, bullet=False)
    notes(
        s,
        """
Em chia 12 tuần thực tập thành bốn giai đoạn.
Giai đoạn một, tuần 1 đến 3: nghiên cứu và phân tích yêu cầu. Em tìm hiểu lĩnh vực fulfillment, khách hàng mục tiêu, khảo sát bố cục và cấu trúc website Printuz, và học Bootstrap. Kết quả là có định hướng thiết kế và danh sách các trang cần xây dựng.
Giai đoạn hai, tuần 4 đến 8: xây dựng giao diện. Em làm phần dùng chung và trang chủ trước, sau đó là catalog, chi tiết sản phẩm, giới thiệu, liên hệ, blog, và cuối cùng là responsive.
Giai đoạn ba, tuần 9 đến 10: hoàn thiện và kiểm thử, sửa lỗi theo góp ý và báo cáo với người hướng dẫn. Kết quả là bản demo hoàn chỉnh.
Giai đoạn bốn, hai tuần cuối, em tổng hợp và viết báo cáo thực tập.
""",
    )

    # ---------------------------------------------------------------- 13 Log table
    s = new("04 · Kế hoạch và tiến độ", "Nhật ký thực tập 12 tuần")
    log = [
        ["Tuần", "Thời gian", "Công việc được giao", "Nội dung chính"],
        ["1", "06/07 – 12/07", "Tìm hiểu website Printuz", "Nắm bắt yêu cầu và định hướng của kỳ thực tập"],
        ["2", "13/07 – 19/07", "Phân tích giao diện", "Khảo sát cấu trúc, các trang chính, danh mục sản phẩm; xác định phạm vi"],
        ["3", "20/07 – 26/07", "Tìm hiểu Bootstrap", "Các thành phần và cách dùng Bootstrap trong thiết kế giao diện"],
        ["4", "27/07 – 02/08", "Giao diện chung, trang chủ", "Header, Navigation, Footer; hoàn thiện trang chủ"],
        ["5", "03/08 – 09/08", "Catalog và Shop", "Danh mục, thẻ sản phẩm; trang chi tiết sản phẩm"],
        ["6", "10/08 – 16/08", "About và Contact", "Trang giới thiệu, liên hệ; nội dung, hình ảnh, biểu mẫu"],
        ["7", "17/08 – 23/08", "Blog và trang nội dung", "Danh sách bài viết, trang hiển thị nội dung bài viết"],
        ["8", "24/08 – 30/08", "Responsive bằng Bootstrap", "Grid, breakpoint cho máy tính, máy tính bảng, điện thoại"],
        ["9", "31/08 – 06/09", "Kiểm tra và sửa lỗi", "Bố cục, kích thước, khoảng cách, font, trình duyệt; sửa theo góp ý"],
        ["10", "07/09 – 13/09", "Báo cáo người hướng dẫn", "Hoàn chỉnh các task đã làm để báo cáo"],
        ["11", "14/09 – 20/09", "Viết báo cáo", "Tổng hợp nội dung, viết báo cáo"],
        ["12", "21/09 – 27/09", "Viết báo cáo", "Hoàn tất báo cáo"],
    ]
    phase_of = [0, 0, 0, 1, 1, 1, 1, 1, 2, 2, 3, 3]
    table(s, 0.5, 1.25, [0.75, 1.65, 3.15, 6.78], log, 13, 13, 0.4, [PHASE_TINT[p] for p in phase_of])
    legend = ["Giai đoạn 1", "Giai đoạn 2", "Giai đoạn 3", "Giai đoạn 4"]
    for i, lab in enumerate(legend):
        rect(s, 0.5 + i * 2.0, 6.7, 0.22, 0.2, PHASE_TINT[i], LINE, MSO_SHAPE.RECTANGLE)
        add_text(s, lab, 0.78 + i * 2.0, 6.63, 1.6, 0.3, 11, False, MUTED)
    notes(
        s,
        """
Đây là nhật ký thực tập chi tiết theo từng tuần, cũng là Bảng 2.1 trong báo cáo. Màu nền các dòng tương ứng với bốn giai đoạn.
Tuần 1 và 2 em tìm hiểu website Printuz và phân tích giao diện. Tuần 3 em học Bootstrap.
Tuần 4 em làm header, menu, footer và trang chủ. Tuần 5 làm catalog và trang chi tiết sản phẩm. Tuần 6 làm trang giới thiệu và liên hệ. Tuần 7 làm blog. Tuần 8 chỉnh responsive.
Tuần 9 em kiểm tra và sửa lỗi toàn bộ các trang. Tuần 10 hoàn chỉnh để báo cáo với người hướng dẫn.
Hai tuần cuối em viết và hoàn tất báo cáo.
""",
    )

    # ---------------------------------------------------------------- 14 Analysis
    s = new("05 · Phân tích và công nghệ", "Phân tích thực trạng theo chủ đề thực tập")
    an = [
        ("GIAO DIỆN NGƯỜI DÙNG (UI)", "Rõ ràng, thống nhất",
         ["Website fulfillment có nhiều nhóm nội dung: doanh nghiệp, dịch vụ, sản phẩm, quy trình, tin tức, FAQ, liên hệ.",
          "Cần bố cục rõ, phân cấp thông tin hợp lý.",
          "Phong cách hiện đại, tối giản, thống nhất màu thương hiệu.",
          "Dùng hình ảnh sản phẩm và khối nội dung trực quan."]),
        ("TRẢI NGHIỆM NGƯỜI DÙNG (UX)", "Dễ tìm, dễ dùng",
         ["Menu sắp theo mức ưu tiên, tên gọi dễ hiểu.",
          "Ô tìm kiếm đặt ở chỗ dễ thấy.",
          "Hiển thị tốt trên nhiều kích thước màn hình.",
          "Hiệu ứng vừa phải: di chuột, nút về đầu trang, menu cố định khi cuộn."]),
        ("YÊU CẦU DỰ ÁN", "Đa trang, đồng bộ",
         ["Website đa trang giới thiệu thương hiệu và dịch vụ.",
          "Giao diện nhất quán giữa các trang.",
          "Thành phần dùng chung cho bố cục, màu sắc, điều hướng.",
          "Tổ chức dữ liệu sản phẩm và bài viết để mở rộng."]),
    ]
    for i, (k, t, items) in enumerate(an):
        card(s, 0.5 + i * 4.18, 1.4, 3.97, 5.35, k, t, items, 16, fill=SOFT if i != 1 else SOFT_2)
    notes(
        s,
        """
Trước khi xây dựng, em phân tích trên ba khía cạnh.
Về giao diện người dùng: website fulfillment có nhiều nhóm nội dung, nên cần bố cục rõ ràng và phân cấp thông tin. Phong cách nên hiện đại, tối giản và thống nhất màu thương hiệu, có hình ảnh sản phẩm và khối nội dung trực quan.
Về trải nghiệm người dùng: menu phải sắp theo mức ưu tiên và dùng tên dễ hiểu, tìm kiếm phải dễ thấy, giao diện phải hiển thị tốt trên điện thoại. Các hiệu ứng như menu cố định hay nút về đầu trang giúp ích, nhưng phải dùng vừa phải để không rối mắt và không làm chậm trang.
Về yêu cầu dự án: website đa trang, nhất quán, có thành phần dùng chung, và dữ liệu sản phẩm, bài viết được tổ chức để dễ mở rộng.
""",
    )

    # ---------------------------------------------------------------- 15 Technology
    s = new("05 · Phân tích và công nghệ", "Kỹ thuật và công nghệ sử dụng")
    tech = [
        ("HTML", "Cấu trúc nội dung", "Tiêu đề, đoạn văn, hình ảnh, menu, nút, biểu mẫu, danh sách sản phẩm."),
        ("CSS", "Định dạng giao diện", "Màu sắc, font, khoảng cách, hiệu ứng di chuột, menu cố định khi cuộn."),
        ("JavaScript", "Tương tác", "Tìm kiếm, ẩn hiện câu trả lời FAQ, menu điện thoại, về đầu trang, hiển thị dữ liệu."),
        ("Bootstrap 5", "Bố cục đáp ứng", "Hệ lưới, breakpoint, nút, biểu mẫu, thanh điều hướng."),
        ("JSON", "Dữ liệu", "Danh sách sản phẩm, danh mục, bài viết; tách dữ liệu khỏi giao diện."),
    ]
    for i, (name, role, body) in enumerate(tech):
        x = 0.5 + i * 2.5
        rect(s, x, 1.4, 2.33, 3.9, SOFT if i % 2 == 0 else SOFT_2)
        add_text(s, name, x + 0.15, 1.55, 2.05, 0.55, 22, True, TEAL_DARK)
        add_text(s, role, x + 0.15, 2.12, 2.05, 0.4, 15, True, INK)
        add_text(s, body, x + 0.15, 2.6, 2.05, 2.6, 16, False, MUTED)
    rect(s, 0.5, 5.55, 12.33, 1.2, WHITE, LINE)
    add_text(s, "NHẬN DIỆN THỐNG NHẤT", 0.75, 5.65, 4, 0.3, 12, True, TEAL_DARK)
    sw = rect(s, 0.78, 6.05, 0.5, 0.5, TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(
        s,
        "Màu chủ đạo xanh teal của Printuz  ·  Tiêu đề dùng font Poppins, nội dung dùng Inter  ·  Header, menu, footer, tìm kiếm dùng chung cho mọi trang",
        1.45,
        6.02,
        11.2,
        0.6,
        15,
        False,
        INK,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    notes(
        s,
        """
Em dùng năm công nghệ chính.
HTML để dựng cấu trúc nội dung: tiêu đề, hình ảnh, menu, nút, biểu mẫu và danh sách sản phẩm.
CSS để định dạng màu, font, khoảng cách và các hiệu ứng như đổi trạng thái khi di chuột hay menu cố định khi cuộn.
JavaScript để tạo tương tác: tìm kiếm, ẩn hiện câu trả lời trong trang câu hỏi thường gặp, menu trên điện thoại, nút về đầu trang và hiển thị dữ liệu.
Bootstrap 5 cung cấp hệ lưới và breakpoint, giúp em làm giao diện đáp ứng nhanh hơn.
JSON dùng để lưu sản phẩm, danh mục và bài viết, tách dữ liệu ra khỏi giao diện, nên sau này dễ nối với cơ sở dữ liệu hoặc API.
Toàn bộ website dùng chung màu teal của Printuz, font Poppins cho tiêu đề và Inter cho nội dung.
""",
    )

    # ---------------------------------------------------------------- 16 Result: home + about
    s = new("06 · Kết quả thực hiện", "Trang chủ và trang giới thiệu doanh nghiệp")
    figure(s, SHOTS / "printuz-home.png", 0.5, 1.35, 6.05, 4.6, "Hình 2.1. Trang chủ website Printuz")
    figure(s, SHOTS / "printuz-about.png", 6.78, 1.35, 6.05, 4.6, "Hình 2.2. Trang giới thiệu doanh nghiệp")
    add_text(
        s,
        "Header, menu và footer dùng chung. Trang giới thiệu trình bày mạng lưới fulfillment Việt Nam, Trung Quốc, Hoa Kỳ, bốn dịch vụ và lý do chọn Printuz.",
        0.5,
        6.3,
        12.33,
        0.65,
        15,
        False,
        INK,
    )
    notes(
        s,
        """
Đây là kết quả của giai đoạn xây dựng giao diện.
Hình 2.1 là trang chủ, có thanh thương hiệu, ô tìm kiếm, menu chính, khối giới thiệu và chân trang. Phần header, menu và footer này dùng chung cho tất cả các trang.
Hình 2.2 là trang giới thiệu doanh nghiệp. Trang này trình bày Printuz là gì, mạng lưới sản xuất ở Việt Nam, Trung Quốc và Hoa Kỳ, bốn dịch vụ chính là Print-on-Demand, dropshipping, xử lý đơn và giải pháp riêng, cùng lý do khách hàng nên chọn Printuz.
""",
    )

    # ---------------------------------------------------------------- 17 Result: catalog + product
    s = new("06 · Kết quả thực hiện", "Danh mục sản phẩm và trang chi tiết sản phẩm")
    figure(s, SHOTS / "printuz-shop.png", 0.5, 1.35, 6.05, 4.6, "Hình 2.3. Giao diện danh mục sản phẩm")
    figure(s, SHOTS / "printuz-product.png", 6.78, 1.35, 6.05, 4.6, "Hình 2.7. Giao diện chi tiết sản phẩm")
    add_text(
        s,
        "634 sản phẩm hiển thị từ dữ liệu JSON. Lọc theo cây danh mục và khoảng giá, sắp xếp, phân trang 28 sản phẩm mỗi trang. Trang chi tiết có ảnh, giá, thông số, sản phẩm liên quan.",
        0.5,
        6.3,
        12.33,
        0.65,
        15,
        False,
        INK,
    )
    notes(
        s,
        """
Hình 2.3 là giao diện danh mục sản phẩm. Bên trái là cây danh mục và bộ lọc giá, bên phải là các thẻ sản phẩm gồm hình ảnh, nhóm sản phẩm, tên và giá.
Toàn bộ catalog có 634 sản phẩm, được hiển thị từ dữ liệu JSON. Trong ảnh đang mở danh mục Print-on-Demand với 611 sản phẩm, mỗi trang 28 sản phẩm.
Hình 2.7 là trang chi tiết sản phẩm. Ví dụ Custom Printed PC Mouse có bộ ảnh, giá 12,99 đô la, mã sản phẩm, thời gian sản xuất, nút bắt đầu đơn, và bên dưới là phần vận chuyển, mô tả, sản phẩm liên quan.
""",
    )

    # ---------------------------------------------------------------- 18 Result: blog + faq
    s = new("06 · Kết quả thực hiện", "Trang tin tức và trang câu hỏi thường gặp")
    figure(s, SHOTS / "printuz-blog.png", 0.5, 1.35, 6.05, 4.6, "Hình 2.5. Giao diện trang tin tức")
    figure(s, SHOTS / "printuz-faq-open.png", 6.78, 1.35, 6.05, 4.6,
           "Hình 2.9. Hiển thị nội dung câu hỏi thường gặp")
    add_text(
        s,
        "Tin tức: 1.092 bài, xem theo danh mục, thẻ, tác giả, có phân trang. FAQ: câu hỏi chia theo nhóm, bấm vào câu hỏi để mở câu trả lời.",
        0.5,
        6.3,
        12.33,
        0.65,
        15,
        False,
        INK,
    )
    notes(
        s,
        """
Hình 2.5 là trang tin tức. Cột trái là danh mục kèm số bài, phần chính là danh sách bài viết dạng thẻ. Trong ảnh là danh mục Printuz Guide gồm 11 bài hướng dẫn cho người bán.
Toàn bộ blog có 1.092 bài, xem được theo danh mục, thẻ hoặc tác giả, và có phân trang.
Hình 2.9 là trang câu hỏi thường gặp. Câu hỏi được chia theo nhóm như giới thiệu Printuz, Print-on-Demand, xử lý đơn, dropshipping, vận chuyển. Người dùng bấm vào câu hỏi thì câu trả lời mở ra, bấm lần nữa thì đóng lại. Phần này em viết bằng JavaScript.
""",
    )

    # ---------------------------------------------------------------- 19 Result: sticky + mobile
    s = new("06 · Kết quả thực hiện", "Thanh điều hướng cố định và giao diện di động")
    figure(s, SHOTS / "printuz-sticky.png", 0.5, 1.35, 7.1, 4.75,
           "Hình 2.6. Thanh điều hướng cố định khi cuộn trang")
    picture(s, SHOTS / "printuz-mobile-about.png", 7.95, 1.35, 2.3, 4.75)
    picture(s, SHOTS / "printuz-mobile-menu.png", 10.5, 1.35, 2.3, 4.75)
    add_text(s, "Hình 2.8. Giao diện trên điện thoại", 7.8, 6.16, 5.1, 0.3, 12, False,
             MUTED, PP_ALIGN.CENTER)
    add_text(
        s,
        "Khi cuộn xuống, menu chính trượt xuống và giữ ở đầu màn hình. Trên điện thoại, menu thu vào nút ba gạch, mở ra kèm ô tìm kiếm.",
        0.5,
        6.48,
        12.33,
        0.5,
        15,
        False,
        INK,
    )
    notes(
        s,
        """
Hình 2.6 là thanh điều hướng cố định. Khi người dùng cuộn xuống quá một đoạn, menu chính trượt xuống và giữ ở đầu màn hình, nên luôn chuyển trang được mà không cần cuộn ngược lên.
Hình 2.8 là giao diện trên điện thoại, kết quả của tuần 8 khi em chỉnh responsive bằng Bootstrap.
Nội dung tự xếp thành một cột. Menu thu gọn vào nút ba gạch, bấm vào thì mở ra danh sách trang kèm ô tìm kiếm.
""",
    )

    # ---------------------------------------------------------------- 20 Delivered products
    s = new("06 · Kết quả thực hiện", "Sản phẩm đã hoàn thành")
    table(
        s,
        0.5,
        1.4,
        [0.8, 2.4, 3.0],
        [["STT", "Tên sản phẩm", "Mô tả"], ["1", "Website Printuz", "printuz-six.vercel.app"]],
        14,
        15,
        0.6,
    )
    add_text(s, "Bảng 2.2. Danh sách sản phẩm đã hoàn thành", 0.5, 2.67, 6.2, 0.3, 12, False, MUTED)
    stat(s, 0.5, 3.15, 1.95, 1.6, "11", "trang HTML")
    stat(s, 2.62, 3.15, 1.95, 1.6, "634", "sản phẩm")
    stat(s, 4.74, 3.15, 1.95, 1.6, "1.092", "bài viết")
    rect(s, 0.5, 4.95, 6.19, 1.8, SOFT_2)
    tf = textbox(s, 0.7, 5.05, 5.8, 1.65)
    write_items(
        tf,
        [
            "Hiển thị được trên máy tính, máy tính bảng và điện thoại.",
            "Thành phần dùng chung: header, menu, footer, tìm kiếm, đăng nhập, menu di động, nút về đầu trang.",
        ],
        14,
        INK,
        5,
    )
    card(
        s,
        7.0,
        1.4,
        5.83,
        5.35,
        "CÁC TRANG ĐÃ XÂY DỰNG",
        None,
        [
            ("Trang chủ", ""),
            ("Giới thiệu doanh nghiệp, ", "kèm các dịch vụ"),
            ("Danh mục và chi tiết sản phẩm", ""),
            ("Tin tức (blog) ", "và trang đọc bài"),
            ("Câu hỏi thường gặp", ""),
            ("Liên hệ", ""),
            ("Chính sách bảo mật", ""),
            ("Tuyển dụng", ""),
            ("Hướng dẫn sử dụng Extension", ""),
            ("Tài khoản: ", "đăng nhập, đặt lại mật khẩu"),
        ],
        15,
    )
    notes(
        s,
        """
Sản phẩm em hoàn thành là website Printuz, đang chạy bản demo tại printuz-six.vercel.app.
Website gồm 11 trang HTML: trang chủ, giới thiệu, danh mục và chi tiết sản phẩm, tin tức, câu hỏi thường gặp, liên hệ, chính sách bảo mật, tuyển dụng, hướng dẫn extension, và hai trang tài khoản là đăng nhập và đặt lại mật khẩu.
Phần sản phẩm hiển thị 634 sản phẩm, phần tin tức hiển thị 1.092 bài viết, đều lấy từ dữ liệu JSON.
Các thành phần dùng chung như header, menu, footer, tìm kiếm, menu di động và nút về đầu trang giúp mọi trang thống nhất. Website hiển thị được trên máy tính, máy tính bảng và điện thoại.
""",
    )

    # ---------------------------------------------------------------- 21 Pros & cons
    s = new("07 · Đánh giá và kiến nghị", "Ưu điểm và hạn chế của website")
    card(
        s,
        0.5,
        1.4,
        6.05,
        5.35,
        "ƯU ĐIỂM",
        "Rõ ràng, chuyên nghiệp, dùng được trên nhiều thiết bị",
        [
            "Đa trang, nội dung phân loại rõ, dễ tiếp cận từng nhóm thông tin.",
            "Phong cách hiện đại, màu thương hiệu và kiểu chữ phù hợp, tăng nhận diện và sự tin cậy.",
            "Đáp ứng máy tính, máy tính bảng, điện thoại.",
            "Menu, tìm kiếm, danh mục sản phẩm, FAQ, liên hệ giúp sử dụng thuận tiện.",
        ],
        17,
    )
    card(
        s,
        6.78,
        1.4,
        6.05,
        5.35,
        "HẠN CHẾ",
        "Mới tập trung vào giao diện và hiển thị",
        [
            "Chưa có chức năng nghiệp vụ: đăng ký tài khoản thật, xử lý form liên hệ, tải tệp thiết kế, thanh toán, quản lý đơn, đồng bộ dữ liệu fulfillment.",
            "Dữ liệu sản phẩm, bài viết cần nối với hệ quản trị nội dung hoặc cơ sở dữ liệu để cập nhật linh hoạt.",
            "Khi vận hành thật cần thêm bảo mật, xác thực người dùng, sao lưu dữ liệu, kiểm thử hiệu năng.",
        ],
        17,
        fill=SOFT_2,
    )
    notes(
        s,
        """
Về ưu điểm, website chia đa trang nên nội dung phân loại rõ ràng. Giao diện hiện đại, dùng màu thương hiệu, giúp tăng nhận diện và sự tin cậy. Website dùng được trên máy tính và điện thoại, có menu, tìm kiếm, danh mục và FAQ giúp người dùng thuận tiện.
Về hạn chế, phiên bản hiện tại chủ yếu là giao diện và hiển thị nội dung. Các chức năng nghiệp vụ như đăng ký tài khoản thật, xử lý form liên hệ, tải tệp thiết kế, thanh toán, quản lý đơn hàng chưa được triển khai.
Dữ liệu cũng cần được nối với hệ quản trị nội dung hoặc cơ sở dữ liệu. Nếu đưa vào vận hành thật thì cần bổ sung bảo mật, xác thực, sao lưu và kiểm thử hiệu năng.
Những hạn chế này cho thấy hướng phát triển ở giai đoạn sau, còn bản hiện tại là nền tảng giao diện.
""",
    )

    # ---------------------------------------------------------------- 22 Challenges
    s = new("07 · Đánh giá và kiến nghị", "Thách thức và giải pháp")
    ch = [
        ["Thách thức", "Giải pháp"],
        ["Giữ giao diện thống nhất giữa nhiều trang và nhiều loại nội dung",
         "Xác định trước màu chủ đạo, font chữ, kiểu nút, cấu trúc điều hướng và cách trình bày; dùng chung CSS và thành phần header, menu, footer"],
        ["Vừa đẹp vừa dễ sử dụng",
         "Chia nội dung thành từng nhóm rõ ràng, dùng tiêu đề và khoảng trắng hợp lý, đặt thông tin quan trọng ở khu vực dễ quan sát"],
        ["Hiển thị tốt trên nhiều thiết bị",
         "Xây dựng theo hướng đáp ứng với Bootstrap Grid và breakpoint; làm menu riêng cho điện thoại"],
        ["Khu vực sản phẩm và bài viết có nhiều dữ liệu",
         "Tách dữ liệu ra file JSON, dùng JavaScript để hiển thị, lọc và phân trang"],
    ]
    table(s, 0.5, 1.4, [4.6, 7.73], ch, 15, 15, 0.92)
    add_text(
        s,
        "Bài học: định hướng thiết kế rõ ngay từ đầu giúp tiết kiệm nhiều thời gian sửa ở giai đoạn sau.",
        0.5,
        6.45,
        12.33,
        0.45,
        15,
        True,
        TEAL_DARK,
    )
    notes(
        s,
        """
Trong quá trình làm, em gặp bốn thách thức chính.
Thứ nhất là giữ giao diện thống nhất khi website có nhiều trang. Em giải quyết bằng cách chốt trước màu, font, kiểu nút, cấu trúc điều hướng, và dùng chung CSS cùng các thành phần header, menu, footer.
Thứ hai là vừa đẹp vừa dễ dùng. Em chia nội dung thành từng nhóm, dùng tiêu đề và khoảng trắng hợp lý, đưa thông tin quan trọng lên chỗ dễ thấy.
Thứ ba là hiển thị trên nhiều thiết bị. Em dùng hệ lưới và breakpoint của Bootstrap, và làm menu riêng cho điện thoại.
Thứ tư là khu vực sản phẩm và bài viết có nhiều dữ liệu. Em tách dữ liệu ra JSON và dùng JavaScript để hiển thị, lọc và phân trang.
Bài học em rút ra là có định hướng thiết kế rõ từ đầu sẽ đỡ sửa rất nhiều về sau.
""",
    )

    # ---------------------------------------------------------------- 23 Future
    s = new("07 · Đánh giá và kiến nghị", "Định hướng phát triển")
    fut = [
        ("Quản trị nội dung", "Tích hợp hệ quản trị nội dung để cập nhật sản phẩm, bài viết, dịch vụ dễ dàng."),
        ("Tài khoản người dùng", "Đăng ký, đăng nhập và quản lý tài khoản."),
        ("Hỗ trợ khách hàng", "Biểu mẫu liên hệ, tải tệp thiết kế, hỗ trợ trực tuyến."),
        ("Đơn hàng và thanh toán", "Kết nối quản lý đơn hàng, thanh toán trực tuyến, dữ liệu fulfillment."),
        ("Tối ưu", "SEO, tốc độ tải trang và bảo mật website."),
        ("Hỗ trợ kinh doanh", "Quản lý sản phẩm, theo dõi đơn hàng, thống kê hoạt động."),
    ]
    for i, (head, body) in enumerate(fut):
        col, row = i % 3, i // 3
        x = 0.5 + col * 4.18
        y = 1.4 + row * 2.72
        rect(s, x, y, 3.97, 2.5, SOFT if (col + row) % 2 == 0 else SOFT_2)
        add_text(s, f"0{i + 1}", x + 0.2, y + 0.15, 1, 0.5, 24, True, TEAL)
        add_text(s, head, x + 0.2, y + 0.7, 3.6, 0.45, 18, True, INK)
        add_text(s, body, x + 0.2, y + 1.18, 3.6, 1.2, 15, False, MUTED)
    notes(
        s,
        """
Từ những hạn chế trên, em đề xuất sáu hướng phát triển.
Một là tích hợp hệ quản trị nội dung để cập nhật sản phẩm, bài viết dễ hơn.
Hai là xây dựng đăng ký, đăng nhập và quản lý tài khoản.
Ba là bổ sung biểu mẫu liên hệ, tải tệp thiết kế và hỗ trợ khách hàng trực tuyến.
Bốn là kết nối quản lý đơn hàng, thanh toán trực tuyến và dữ liệu fulfillment.
Năm là tối ưu SEO, tốc độ tải trang và bảo mật.
Sáu là phát triển các tính năng hỗ trợ kinh doanh như quản lý sản phẩm, theo dõi đơn hàng và thống kê.
""",
    )

    # ---------------------------------------------------------------- 24 Self-assessment
    s = new("07 · Đánh giá và kiến nghị", "Nhận xét về đơn vị và bài học của bản thân")
    card(s, 0.5, 1.4, 6.05, 2.55, "VỀ ĐƠN VỊ THỰC TẬP", None,
         ["Môi trường làm việc hiện đại, thân thiện, chuyên nghiệp; mọi người sẵn sàng hỗ trợ, góp ý.",
          "Đội ngũ có trách nhiệm, chủ động, thích ứng nhanh với công nghệ mới; tạo điều kiện cho thực tập sinh học hỏi."],
         16)
    card(s, 6.78, 1.4, 6.05, 2.55, "VỀ CHUYÊN MÔN", None,
         ["Có kinh nghiệm xây dựng giao diện, tổ chức bố cục, cải thiện trải nghiệm người dùng.",
          "Hiểu website cần vừa thẩm mỹ, vừa rõ ràng, dễ dùng, thống nhất thương hiệu; biết dùng công cụ AI để tự học."],
         16, fill=SOFT_2)
    card(s, 0.5, 4.15, 6.05, 2.6, "VỀ KỸ NĂNG MỀM", None,
         ["Cải thiện làm việc nhóm, giao tiếp, tiếp nhận ý kiến.",
          "Hiểu tầm quan trọng của trao đổi rõ ràng và điều chỉnh sản phẩm theo phản hồi của người hướng dẫn."],
         16, fill=SOFT_2)
    card(s, 6.78, 4.15, 6.05, 2.6, "ĐIỂM CẦN CẢI THIỆN", None,
         ["Quản lý thời gian và dự đoán khối lượng công việc còn hạn chế.",
          "Giai đoạn đầu chưa sắp xếp hợp lý khi nhiều việc cùng lúc; cần lập kế hoạch chủ động hơn."],
         16)
    notes(
        s,
        """
Về đơn vị thực tập, em thấy TGSOFT có môi trường làm việc hiện đại, thân thiện và chuyên nghiệp. Mọi người luôn sẵn sàng hỗ trợ và góp ý, và công ty tạo điều kiện để thực tập sinh học hỏi quy trình thực tế.
Về chuyên môn, em có thêm kinh nghiệm xây dựng giao diện, tổ chức bố cục và cải thiện trải nghiệm người dùng. Em hiểu rằng website không chỉ cần đẹp mà còn phải rõ ràng, dễ dùng và thống nhất với thương hiệu. Các công cụ AI cũng giúp em tự học và tự tin hơn.
Về kỹ năng mềm, em cải thiện khả năng làm việc nhóm, giao tiếp và tiếp nhận góp ý.
Điểm em cần cải thiện là quản lý thời gian và ước lượng khối lượng công việc. Giai đoạn đầu em chưa sắp xếp tốt khi có nhiều việc cùng lúc, và đây là bài học để em lập kế hoạch chủ động hơn.
""",
    )

    # ---------------------------------------------------------------- 25 Recommendations
    s = new("07 · Đánh giá và kiến nghị", "Kiến nghị đối với công ty")
    recs = [
        ("Nâng cấp giao diện và trải nghiệm",
         "Tiếp tục đầu tư giao diện, trải nghiệm người dùng cho các sản phẩm website, giúp nâng cao hình ảnh thương hiệu và hỗ trợ khách hàng tiếp cận thông tin."),
        ("Cho thực tập sinh tham gia dự án thực tế",
         "Giao nhiệm vụ phù hợp năng lực để sinh viên nâng cao chuyên môn, hình thành trách nhiệm và sớm thích nghi với doanh nghiệp."),
        ("Góp ý cụ thể và tận dụng công cụ AI",
         "Góp ý cụ thể từ người hướng dẫn giúp thực tập sinh nhận ra điểm mạnh, điểm cần cải thiện nhanh hơn; công cụ AI hỗ trợ tốt cho việc học và làm."),
    ]
    for i, (head, body) in enumerate(recs):
        y = 1.4 + i * 1.8
        rect(s, 0.5, y, 12.33, 1.6, SOFT if i % 2 == 0 else SOFT_2)
        rect(s, 0.75, y + 0.38, 0.85, 0.85, TEAL, shape=MSO_SHAPE.OVAL)
        add_text(s, str(i + 1), 0.75, y + 0.38, 0.85, 0.85, 24, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(s, head, 1.9, y + 0.17, 10.7, 0.45, 19, True, INK)
        add_text(s, body, 1.9, y + 0.65, 10.7, 0.9, 15, False, MUTED)
    notes(
        s,
        """
Từ trải nghiệm thực tập, em có ba kiến nghị với công ty.
Thứ nhất, công ty có thể tiếp tục đầu tư nâng cấp giao diện và trải nghiệm người dùng cho các sản phẩm website, vì giao diện tốt giúp nâng cao hình ảnh thương hiệu và hỗ trợ khách hàng.
Thứ hai, công ty có thể tiếp tục cho thực tập sinh tham gia nhiều hơn vào dự án thực tế, với nhiệm vụ phù hợp năng lực.
Thứ ba, những góp ý cụ thể từ người hướng dẫn giúp thực tập sinh tiến bộ nhanh, và các công cụ AI cũng có thể được tận dụng nhiều hơn trong học tập và làm việc.
""",
    )

    # ---------------------------------------------------------------- 26 Conclusion
    s = new("08 · Kết luận", "Kết luận")
    tf = textbox(s, 0.5, 1.4, 7.4, 4.3)
    write_items(
        tf,
        [
            "Hoàn thành giao diện website giới thiệu dịch vụ cho Printuz LTD, đúng kế hoạch 12 tuần tại TGSOFT.",
            "Vận dụng HTML, CSS, JavaScript, Bootstrap và JSON vào một dự án thực tế.",
            "Nắm được quy trình làm việc: nhận yêu cầu, xây dựng, kiểm tra, sửa lỗi, báo cáo.",
            "Rèn luyện tự học, làm việc nhóm, giao tiếp và tiếp nhận phản hồi.",
            "Website hiện là nền tảng giao diện để phát triển tiếp các chức năng nghiệp vụ.",
        ],
        17,
        INK,
        10,
    )
    rect(s, 8.2, 1.4, 4.63, 4.3, TEAL)
    add_text(s, "LỜI CẢM ƠN", 8.45, 1.6, 4.2, 0.3, 12, True, WHITE)
    add_text(
        s,
        "Em xin chân thành cảm ơn Công ty TGSOFT, anh Trần Thế Văn, anh Trần Thế Phú, cô Nguyễn Thị Định và quý thầy cô Khoa Công nghệ Thông tin.",
        8.45,
        1.95,
        4.15,
        2.4,
        17,
        False,
        WHITE,
    )
    add_text(s, "Em xin lắng nghe câu hỏi và góp ý của quý thầy cô.", 8.45, 4.55, 4.15, 1.0, 17, True, WHITE)
    add_text(s, "Bản demo: printuz-six.vercel.app", 0.5, 6.2, 7.4, 0.4, 15, True, TEAL_DARK)
    notes(
        s,
        """
Tóm lại, trong 12 tuần thực tập tại TGSOFT, em đã hoàn thành giao diện website giới thiệu dịch vụ cho Printuz LTD theo đúng kế hoạch.
Em vận dụng được HTML, CSS, JavaScript, Bootstrap và JSON vào một dự án thực tế, nắm được quy trình làm việc từ nhận yêu cầu đến báo cáo, và rèn thêm kỹ năng tự học, làm việc nhóm, tiếp nhận phản hồi.
Website hiện tại là nền tảng giao diện để phát triển tiếp các chức năng nghiệp vụ.
Em xin chân thành cảm ơn Công ty TGSOFT, anh Trần Thế Văn, anh Trần Thế Phú, cô Nguyễn Thị Định và quý thầy cô trong khoa.
Em xin lắng nghe câu hỏi và góp ý của quý thầy cô.

--- Nếu được hỏi thêm ---
Vì sao chọn website tĩnh, chưa có backend? Phạm vi đề tài là giao diện. Dữ liệu tách ra JSON để sau này thay bằng API hoặc cơ sở dữ liệu mà không phải làm lại giao diện.
Responsive làm thế nào? Em dùng hệ lưới 12 cột và các breakpoint của Bootstrap, kiểm tra ở máy tính, máy tính bảng và điện thoại; trên điện thoại menu thu vào nút ba gạch.
Đăng nhập đã dùng được chưa? Form đã kiểm tra dữ liệu nhập. Xác thực tài khoản thật thuộc ứng dụng app.printuz.com, nằm ngoài phạm vi đề tài.
Em dùng AI vào việc gì? Em dùng để tra cứu, giải thích lỗi và gợi ý cách làm; em tự kiểm tra và chỉnh lại theo góp ý của người hướng dẫn.
""",
    )

    total = len(prs.slides)
    for i, sl in enumerate(prs.slides, 1):
        if i == 1:
            continue
        add_text(sl, f"{i} / {total}", 11.3, 7.1, 1.53, 0.3, 11, False, MUTED, PP_ALIGN.RIGHT)

    prs.save(OUT)
    print(OUT, total)


if __name__ == "__main__":
    build()
